# ============================================================
# PARSER
# ============================================================
# Responsável pela extração do conteúdo textual de arquivos PDF e DOCX.
# Para DOCX, são considerados tanto parágrafos quanto tabelas.

import fitz
from docx import Document

# Extrai o texto página a página de um arquivo PDF.
def extract_pdf(file_path):
    document = fitz.open(file_path)

    text = []

    for page in document:
        text.append(page.get_text())

    document.close()

    return "\n".join(text)

 # Extrai parágrafos e conteúdo de tabelas de um arquivo DOCX.
def extract_docx(file_path):
    document = Document(file_path)

    text = []

    # Extrai os parágrafos
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text.strip())

    # Extrai as tabelas
    for table in document.tables:
        for row in table.rows:
            cells = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    cells.append(cell_text)

            if cells:
                text.append(" | ".join(cells))

    return "\n".join(text)

# Direciona a extração de acordo com a extensão do arquivo.
def extract_text(file_path):
    if file_path.lower().endswith(".pdf"):
        return extract_pdf(file_path)

    elif file_path.lower().endswith(".docx"):
        return extract_docx(file_path)

    else:
        raise ValueError("Formato de arquivo não suportado.")


if __name__ == "__main__":
    caminho = "data/uploads/Empresa 4.docx"

    texto = extract_text(caminho)

    print(texto)